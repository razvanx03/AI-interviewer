import io
import zipfile
import logging
from typing import Tuple, List
import pypdf
import pypdfium2
import docx
from .ocr_service import ocr_service

logger = logging.getLogger("api.services.extractor")

class DocumentExtractor:
    """
    Independent document extraction engine for PDF and DOCX files.
    Performs binary magic-byte validation, in-memory text extraction,
    and automatic per-page/embedded-image PaddleOCR fallback without temporary disk writes.
    """

    SPARSE_TEXT_THRESHOLD: int = 50  # Character count below which OCR fallback is triggered
    MAX_FILE_SIZE_BYTES: int = 10 * 1024 * 1024  # 10 MB maximum upload limit per file

    @staticmethod
    def validate_file(filename: str, file_bytes: bytes) -> str:
        """
        Validate file extension, file size, and binary magic bytes.
        Returns validated format: 'pdf' or 'docx'.
        Raises ValueError with clear message on invalid, empty, or oversized files.
        """
        if not file_bytes or len(file_bytes) == 0:
            raise ValueError("Uploaded file is empty (0 bytes).")

        if len(file_bytes) > DocumentExtractor.MAX_FILE_SIZE_BYTES:
            raise ValueError(f"Uploaded file '{filename}' exceeds maximum allowed size of 10MB.")

        fname_lower = filename.lower()
        if not (fname_lower.endswith(".pdf") or fname_lower.endswith(".docx")):
            raise ValueError("Unsupported file format. Only .pdf and .docx documents are accepted.")

        # Magic Bytes Validation
        if fname_lower.endswith(".pdf"):
            if not file_bytes.startswith(b"%PDF-"):
                raise ValueError("Corrupted or invalid PDF document: Missing standard PDF file header.")
            return "pdf"

        if fname_lower.endswith(".docx"):
            # DOCX must start with standard ZIP local file header PK\x03\x04
            if not file_bytes.startswith(b"PK\x03\x04"):
                raise ValueError("Corrupted or invalid DOCX document: Not a valid OpenXML package.")
            try:
                with zipfile.ZipFile(io.BytesIO(file_bytes)) as zf:
                    namelist = zf.namelist()
                    if "[Content_Types].xml" not in namelist and "word/document.xml" not in namelist:
                        raise ValueError("Corrupted DOCX document: Missing Word document structure.")
            except zipfile.BadZipFile as exc:
                raise ValueError(f"Corrupted DOCX archive: {exc}") from exc
            return "docx"

        raise ValueError("Unsupported document format.")

    def extract_docx(self, file_bytes: bytes) -> str:
        """
        Extract text from DOCX paragraphs and tables.
        If extracted text is sparse, extracts embedded images from the ZIP stream and runs OCR.
        """
        text_lines: List[str] = []
        try:
            doc = docx.Document(io.BytesIO(file_bytes))
            for p in doc.paragraphs:
                p_text = p.text.strip()
                if p_text:
                    text_lines.append(p_text)

            for table in doc.tables:
                for row in table.rows:
                    row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if row_cells:
                        text_lines.append(" | ".join(row_cells))
        except Exception as exc:
            logger.error("python-docx extraction failed: %s", exc, exc_info=True)
            raise ValueError(f"Failed to parse Word document: {exc}") from exc

        combined_text = "\n".join(text_lines).strip()

        # Check if text is sparse and DOCX contains embedded images
        if len(combined_text) < self.SPARSE_TEXT_THRESHOLD:
            logger.info("DOCX text is sparse (%d chars). Searching for embedded images for OCR...", len(combined_text))
            ocr_text_parts: List[str] = []
            try:
                with zipfile.ZipFile(io.BytesIO(file_bytes)) as zf:
                    image_filenames = [
                        f for f in zf.namelist()
                        if f.startswith("word/media/") and f.lower().endswith((".png", ".jpg", ".jpeg", ".bmp", ".tiff", ".webp"))
                    ]
                    for img_name in image_filenames:
                        img_bytes = zf.read(img_name)
                        img_ocr_text = ocr_service.extract_text_from_image(img_bytes)
                        if img_ocr_text.strip():
                            ocr_text_parts.append(img_ocr_text.strip())
            except Exception as ocr_exc:
                logger.warning("DOCX embedded image OCR extraction encountered an error: %s", ocr_exc)

            if ocr_text_parts:
                combined_text = (combined_text + "\n\n" + "\n\n".join(ocr_text_parts)).strip()

        return combined_text

    def extract_pdf(self, file_bytes: bytes) -> str:
        """
        Extract text from PDF page-by-page using pypdf.
        If an individual page has sparse/no text (e.g. scanned), renders it in-memory via pypdfium2 and runs OCR.
        """
        try:
            pdf_reader = pypdf.PdfReader(io.BytesIO(file_bytes))
        except Exception as exc:
            logger.error("pypdf initialization failed: %s", exc, exc_info=True)
            raise ValueError(f"Failed to read PDF document structure: {exc}") from exc

        if pdf_reader.is_encrypted:
            try:
                # Try decrypting with empty password
                pdf_reader.decrypt("")
            except Exception as exc:
                raise ValueError("PDF is password-protected and cannot be read.") from exc

        page_texts: List[str] = []
        pdfium_doc = None  # Lazy-load only if OCR page rendering is needed

        for i, page in enumerate(pdf_reader.pages):
            extracted_page_text = ""
            try:
                extracted_page_text = (page.extract_text() or "").strip()
            except Exception as page_exc:
                logger.warning("pypdf failed on page %d: %s", i + 1, page_exc)

            # If page text is sufficient, keep it
            if len(extracted_page_text) >= self.SPARSE_TEXT_THRESHOLD:
                page_texts.append(extracted_page_text)
            else:
                # Run OCR on sparse / scanned page
                logger.info("PDF page %d has sparse text (%d chars). Running in-memory OCR...", i + 1, len(extracted_page_text))
                ocr_page_text = ""
                try:
                    if pdfium_doc is None:
                        pdfium_doc = pypdfium2.PdfDocument(file_bytes)
                    
                    pdfium_page = pdfium_doc[i]
                    pil_img = pdfium_page.render(scale=2.0).to_pil()
                    ocr_page_text = ocr_service.extract_text_from_image(pil_img).strip()
                except Exception as ocr_exc:
                    logger.warning("OCR rendering failed on PDF page %d: %s", i + 1, ocr_exc)

                # Combine pypdf text and OCR text if both have content
                final_page_text = "\n".join(filter(None, [extracted_page_text, ocr_page_text])).strip()
                if final_page_text:
                    page_texts.append(final_page_text)

        return "\n\n".join(page_texts).strip()

    def extract_text(self, filename: str, file_bytes: bytes) -> Tuple[str, str]:
        """
        Main extraction entry point.
        Validates document and returns (raw_text, file_type).
        """
        file_type = self.validate_file(filename, file_bytes)

        if file_type == "pdf":
            raw_text = self.extract_pdf(file_bytes)
        elif file_type == "docx":
            raw_text = self.extract_docx(file_bytes)
        else:
            raise ValueError(f"Unsupported file type: {file_type}")

        # Sanitize text: remove null bytes (\x00) and orphaned control chars that break PostgreSQL text encoding
        from core.constants import sanitize_postgres_text
        raw_text = sanitize_postgres_text(raw_text)

        if not raw_text or len(raw_text.strip()) < 10:
            raise ValueError("No readable text could be extracted from document. The file may be empty or contain unsupported image formats.")

        return raw_text, file_type

    def extract_as_document(self, filename: str, file_bytes: bytes):
        """
        Extracts document text and wraps it in a standardized LangChain Document object
        with complete metadata (source, file_type, char_count).
        """
        from langchain_core.documents import Document
        raw_text, file_type = self.extract_text(filename, file_bytes)
        return Document(
            page_content=raw_text,
            metadata={
                "source": filename,
                "file_type": file_type,
                "char_count": len(raw_text),
            },
        )

document_extractor = DocumentExtractor()

