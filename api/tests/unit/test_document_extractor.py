import io
import zipfile
import pytest
import pypdf
import docx

from services.document_extractor import DocumentExtractor, document_extractor


class TestDocumentExtractorUnit:
    """Unit tests for DocumentExtractor validation, DOCX parsing, PDF handling, and text sanitization."""

    def test_validate_file_valid_pdf(self):
        fake_pdf = b"%PDF-1.4 sample pdf binary data"
        fmt = DocumentExtractor.validate_file("candidate.pdf", fake_pdf)
        assert fmt == "pdf"

    def test_validate_file_valid_docx(self):
        # Create a valid minimal DOCX zip in-memory
        bio = io.BytesIO()
        with zipfile.ZipFile(bio, "w") as zf:
            zf.writestr("[Content_Types].xml", "<Types></Types>")
            zf.writestr("word/document.xml", "<w:document></w:document>")
        docx_bytes = bio.getvalue()

        fmt = DocumentExtractor.validate_file("resume.docx", docx_bytes)
        assert fmt == "docx"

    def test_validate_file_empty_bytes(self):
        with pytest.raises(ValueError, match="Uploaded file is empty"):
            DocumentExtractor.validate_file("cv.pdf", b"")

    def test_validate_file_oversized_raises_error(self):
        oversized = b"a" * (DocumentExtractor.MAX_FILE_SIZE_BYTES + 1)
        with pytest.raises(ValueError, match="exceeds maximum allowed size of 10MB"):
            DocumentExtractor.validate_file("large.pdf", oversized)

    def test_validate_file_unsupported_extension(self):
        with pytest.raises(ValueError, match="Unsupported file format"):
            DocumentExtractor.validate_file("cv.txt", b"Some text")

        with pytest.raises(ValueError, match="Unsupported file format"):
            DocumentExtractor.validate_file("cv.png", b"\x89PNG\r\n\x1a\n")

    def test_validate_pdf_corrupted_header(self):
        with pytest.raises(ValueError, match="Missing standard PDF file header"):
            DocumentExtractor.validate_file("bad.pdf", b"NOT_A_PDF_DATA")

    def test_validate_docx_corrupted_zip(self):
        with pytest.raises(ValueError, match="Not a valid OpenXML package"):
            DocumentExtractor.validate_file("corrupted.docx", b"NOT_A_ZIP_HEADER")

    def test_validate_docx_missing_document_structure(self):
        bio = io.BytesIO()
        with zipfile.ZipFile(bio, "w") as zf:
            zf.writestr("some_random_file.txt", "hello")
        docx_bytes = bio.getvalue()

        with pytest.raises(ValueError, match="Missing Word document structure"):
            DocumentExtractor.validate_file("empty_structure.docx", docx_bytes)

    def test_extract_docx_paragraphs_and_tables(self):
        doc = docx.Document()
        doc.add_paragraph("Alice Smith - Senior Engineer")
        doc.add_paragraph("Skills: Python, FastAPI, Docker, PostgreSQL")

        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Experience"
        table.cell(0, 1).text = "5 Years"
        table.cell(1, 0).text = "Role"
        table.cell(1, 1).text = "Tech Lead"

        bio = io.BytesIO()
        doc.save(bio)
        docx_bytes = bio.getvalue()

        extracted_text = document_extractor.extract_docx(docx_bytes)
        assert "Alice Smith - Senior Engineer" in extracted_text
        assert "Skills: Python, FastAPI, Docker, PostgreSQL" in extracted_text
        assert "Experience | 5 Years" in extracted_text
        assert "Role | Tech Lead" in extracted_text

    def test_extract_text_docx_end_to_end(self):
        doc = docx.Document()
        doc.add_paragraph("Candidate Resume with comprehensive experience in software development.")
        doc.add_paragraph("Key competencies: C#, .NET Core, React.js, Docker containers.")

        bio = io.BytesIO()
        doc.save(bio)
        docx_bytes = bio.getvalue()

        raw_text, file_type = document_extractor.extract_text("candidate_cv.docx", docx_bytes)
        assert file_type == "docx"
        assert "Candidate Resume with comprehensive experience" in raw_text
        assert "C#, .NET Core, React.js, Docker containers." in raw_text

    def test_extract_text_too_sparse_raises_error(self):
        # A document with fewer than 10 characters should raise ValueError
        doc = docx.Document()
        doc.add_paragraph("Hi")
        bio = io.BytesIO()
        doc.save(bio)
        docx_bytes = bio.getvalue()

        with pytest.raises(ValueError, match="No readable text could be extracted"):
            document_extractor.extract_text("sparse.docx", docx_bytes)

